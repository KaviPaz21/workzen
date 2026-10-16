from typing import Literal

from docx import Document
from pydantic import BaseModel, Field

from app.actions.errors import ActionError
from app.actions.paths import free_path, resolve_in_workspace
from app.actions.system import launch
from app.config import Settings

MAX_BLOCKS = 400
MAX_TEXT_CHARACTERS = 20_000
MAX_TABLE_ROWS = 60
MAX_TABLE_COLUMNS = 10


class Block(BaseModel):
    type: Literal["heading", "paragraph", "bullet", "numbered", "table"] = "paragraph"
    text: str = Field(default="", max_length=MAX_TEXT_CHARACTERS)
    rows: list[list[str]] = Field(default_factory=list, max_length=MAX_TABLE_ROWS)


class WordArgs(BaseModel):
    filename: str
    title: str = ""
    blocks: list[Block] = Field(max_length=MAX_BLOCKS)
    open_after: bool = True


STYLES = {"bullet": "List Bullet", "numbered": "List Number"}


def add_table(document, block: Block) -> None:
    if not block.rows or not block.rows[0]:
        raise ActionError("A table block needs rows.")
    columns = len(block.rows[0])
    if columns > MAX_TABLE_COLUMNS or any(len(row) != columns for row in block.rows):
        raise ActionError(f"Every table row must have the same number of cells, at most {MAX_TABLE_COLUMNS}.")
    table = document.add_table(rows=len(block.rows), cols=columns)
    table.style = "Table Grid"
    for number, row in enumerate(block.rows):
        for column, value in enumerate(row):
            cell = table.cell(number, column)
            cell.text = value
            if number == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True


async def create_word_document(settings: Settings, args: dict) -> dict:
    parsed = WordArgs(**args)
    if not parsed.blocks and not parsed.title:
        raise ActionError("The document has no content.")
    name = parsed.filename if parsed.filename.lower().endswith(".docx") else f"{parsed.filename}.docx"
    target = free_path(resolve_in_workspace(settings, name))
    document = Document()
    if parsed.title:
        document.add_heading(parsed.title, level=0)
    for block in parsed.blocks:
        if block.type == "table":
            add_table(document, block)
        elif block.type == "heading":
            document.add_heading(block.text, level=1)
        elif block.type in STYLES:
            document.add_paragraph(block.text, style=STYLES[block.type])
        else:
            document.add_paragraph(block.text)
    document.save(target)
    opened = False
    if parsed.open_after and settings.coworker_launch_apps:
        launch(str(target))
        opened = True
    return {"path": str(target), "filename": target.name, "blocks": len(parsed.blocks), "opened": opened}
