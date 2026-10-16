from urllib.parse import parse_qs, unquote, urlparse

import pytest
from docx import Document
from openpyxl import load_workbook

from app.actions import mail

CASES = {
    "filename": "Cases",
    "sheets": [
        {
            "name": "Test cases",
            "columns": ["Test ID", "Title", "Actual result", "Status"],
            "rows": [["TC-01", "Login", "", "Not run"], ["TC-02", "Logout", "", "Not run"]],
        }
    ],
}


def act(client, auth, name, args):
    return client.post(f"/actions/{name}", json={"args": args}, headers=auth).json()


def test_formulas_are_written_only_when_asked_and_only_when_safe(client, auth):
    sheet = {"name": "Budget", "columns": ["Item", "Cost"], "rows": [["Rent", 1000], ["Food", 400.5], ["Total", "=SUM(B2:B3)"]], "allow_formulas": True, "number_formats": {"Cost": "#,##0.00"}}
    made = act(client, auth, "create_excel_document", {"filename": "Budget", "sheets": [sheet]})
    assert made["ok"], made
    cells = load_workbook(made["result"]["path"])["Budget"]
    assert cells["B4"].value == "=SUM(B2:B3)" and cells["B4"].data_type == "f" and cells["B2"].number_format == "#,##0.00"

    as_text = act(client, auth, "create_excel_document", {"filename": "Text", "sheets": [{**sheet, "allow_formulas": False}]})
    assert load_workbook(as_text["result"]["path"])["Budget"]["B4"].data_type == "s"

    for dangerous in ['=HYPERLINK("http://evil","x")', "=cmd|' /C calc'!A0", "='C:\\secret.xlsx'!A1", "=WEBSERVICE(\"http://x\")", "=[1]Sheet1!A1"]:
        refused = act(client, auth, "create_excel_document", {"filename": "Bad", "sheets": [{**sheet, "rows": [["x", dangerous]]}]})
        assert refused["ok"] is False and "not allowed" in refused["error"], dangerous


def test_results_are_written_back_into_the_same_workbook(client, auth):
    path = act(client, auth, "create_excel_document", CASES)["result"]["path"]
    updated = act(
        client,
        auth,
        "update_excel_document",
        {
            "filename": "Cases.xlsx",
            "updates": [{"sheet": "Test cases", "match_column": "Test ID", "match_value": "TC-01", "set": {"Actual result": "Logged in", "Status": "Pass"}}, {"sheet": "Test cases", "match_column": "Test ID", "match_value": "TC-02", "set": {"Status": "Fail"}}],
            "append_rows": [{"sheet": "Test cases", "rows": [["TC-03", "Extra", "Seen", "Pass"]]}],
            "add_sheets": [{"name": "Summary", "columns": ["Metric", "Value"], "rows": [["Passed", 2], ["Failed", 1]]}],
        },
    )
    assert updated["ok"], updated
    assert updated["result"]["cells_changed"] == 3 and updated["result"]["rows_added"] == 1 and updated["result"]["sheets_added"] == 1
    workbook = load_workbook(path)
    sheet = workbook["Test cases"]
    assert [sheet.cell(row=2, column=column).value for column in (3, 4)] == ["Logged in", "Pass"]
    assert sheet["D3"].value == "Fail" and sheet["A4"].value == "TC-03" and workbook["Summary"]["B2"].value == 2
    assert [file.name for file in __import__("pathlib").Path(path).parent.iterdir()] == ["Cases.xlsx"]


def test_updates_explain_what_is_wrong(client, auth):
    act(client, auth, "create_excel_document", CASES)
    base = {"filename": "Cases.xlsx"}
    wrong_sheet = act(client, auth, "update_excel_document", {**base, "updates": [{"sheet": "Nope", "match_column": "Test ID", "match_value": "TC-01", "set": {"Status": "Pass"}}]})
    assert "no sheet called 'Nope'" in wrong_sheet["error"]
    wrong_row = act(client, auth, "update_excel_document", {**base, "updates": [{"sheet": "Test cases", "match_column": "Test ID", "match_value": "TC-99", "set": {"Status": "Pass"}}]})
    assert "No row" in wrong_row["error"]
    wrong_column = act(client, auth, "update_excel_document", {**base, "updates": [{"sheet": "Test cases", "match_column": "Test ID", "match_value": "TC-01", "set": {"Result": "Pass"}}]})
    assert "no column called 'Result'" in wrong_column["error"]
    assert "not in the workspace" in act(client, auth, "update_excel_document", {"filename": "Missing.xlsx", "updates": []})["error"]
    assert act(client, auth, "update_excel_document", {"filename": "..\\Cases.xlsx"})["ok"] is False


def test_word_documents_can_hold_tables(client, auth):
    blocks = [{"type": "heading", "text": "Plan"}, {"type": "table", "rows": [["Task", "Owner"], ["Design", "Ana"], ["Build", "Raj"]]}, {"type": "paragraph", "text": "Done."}]
    made = act(client, auth, "create_word_document", {"filename": "Plan.docx", "blocks": blocks, "open_after": False})
    assert made["ok"], made
    document = Document(made["result"]["path"])
    assert len(document.tables) == 1 and document.tables[0].cell(2, 1).text == "Raj" and document.tables[0].cell(0, 0).paragraphs[0].runs[0].bold
    ragged = act(client, auth, "create_word_document", {"filename": "Bad.docx", "blocks": [{"type": "table", "rows": [["a", "b"], ["c"]]}], "open_after": False})
    assert ragged["ok"] is False and "same number of cells" in ragged["error"]


def test_a_mail_draft_opens_in_the_mail_app_and_is_never_sent(client, auth, settings, monkeypatch):
    opened = []
    monkeypatch.setattr(mail.os, "startfile", lambda link: opened.append(link), raising=False)
    args = {"to": ["sam@example.com"], "cc": ["lee@example.com"], "subject": "Budget review & next steps", "body": "Hi Sam,\n\nCan we meet on Friday?\n\nThanks,\nPat"}
    assert "switched off" in act(client, auth, "draft_email", args)["error"]
    settings.coworker_launch_apps = True
    result = act(client, auth, "draft_email", args)
    assert result["ok"], result
    assert result["result"]["opened"] is True and "not been sent" in result["result"]["note"]
    link = urlparse(opened[0])
    assert link.scheme == "mailto" and link.path == "sam@example.com"
    query = parse_qs(link.query)
    assert query["subject"] == ["Budget review & next steps"] and unquote(query["body"][0]).startswith("Hi Sam,\r\n\r\nCan we meet") and query["cc"] == ["lee@example.com"]


@pytest.mark.parametrize("address", ["not-an-address", "a@b.com;bcc=boss@x.com", "x@y.com?subject=hi", "a b@c.com"])
def test_bad_addresses_are_refused(client, auth, settings, address):
    settings.coworker_launch_apps = True
    refused = act(client, auth, "draft_email", {"to": [address], "subject": "s", "body": "b"})
    assert refused["ok"] is False and "valid email address" in refused["error"]


def test_a_very_long_mail_is_refused_instead_of_being_cut(client, auth, settings):
    settings.coworker_launch_apps = True
    refused = act(client, auth, "draft_email", {"to": [], "subject": "s", "body": "word " * 1000})
    assert refused["ok"] is False and "too long" in refused["error"]